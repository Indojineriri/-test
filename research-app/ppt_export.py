from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from models import Case
import scraper

TEMPLATE_PATH = Path(__file__).parent / "templates" / "case_template.pptx"

# Shape names in the template (Japanese, as authored). Keep these stable.
SHAPE_TITLE = "タイトル 2"
SHAPE_SUBTITLE = "テキスト ボックス 5"
SHAPE_OVERVIEW = "正方形/長方形 117"
SHAPE_CHALLENGES = "正方形/長方形 118"
SHAPE_SOLUTIONS = "正方形/長方形 119"
SHAPE_PICTURE = "図 10"
SHAPE_LINK = "テキスト ボックス 12"

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _find_shape(slide, name: str):
    for shp in slide.shapes:
        if shp.name == name:
            return shp
    return None


def _set_text_preserving_format(shape, lines: list[str]) -> None:
    """Rewrite the text frame to `lines`, copying the first run's formatting.

    `lines` becomes one paragraph per line. The original text frame's first
    paragraph/run carries font, size and color we want to preserve.
    """
    tf = shape.text_frame
    # Capture template formatting from the first run, if any.
    template_para = tf.paragraphs[0] if tf.paragraphs else None
    template_run = template_para.runs[0] if template_para and template_para.runs else None
    template_bullet_xml = None
    if template_para is not None:
        pPr = template_para._pPr
        if pPr is not None:
            template_bullet_xml = deepcopy(pPr)

    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if template_bullet_xml is not None:
            # Replace the new paragraph's pPr with the template's, preserving bullets.
            existing_pPr = p._pPr
            new_pPr = deepcopy(template_bullet_xml)
            if existing_pPr is not None:
                p._p.remove(existing_pPr)
            p._p.insert(0, new_pPr)
        run = p.add_run()
        run.text = line
        if template_run is not None:
            src_font = template_run.font
            if src_font.size is not None:
                run.font.size = src_font.size
            if src_font.name is not None:
                run.font.name = src_font.name
            if src_font.bold is not None:
                run.font.bold = src_font.bold
            try:
                if src_font.color and src_font.color.rgb is not None:
                    run.font.color.rgb = src_font.color.rgb
            except (AttributeError, TypeError):
                pass


def _set_single_text(shape, text: str) -> None:
    tf = shape.text_frame
    template_run = tf.paragraphs[0].runs[0] if tf.paragraphs and tf.paragraphs[0].runs else None
    src_size = template_run.font.size if template_run else None
    src_name = template_run.font.name if template_run else None
    src_bold = template_run.font.bold if template_run else None
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    if src_size is not None:
        run.font.size = src_size
    if src_name is not None:
        run.font.name = src_name
    if src_bold is not None:
        run.font.bold = src_bold


def _set_link_textbox(slide, shape, link_text: str, url: str) -> None:
    """Set the link textbox content + add a hyperlink relationship on the run."""
    tf = shape.text_frame
    template_run = tf.paragraphs[0].runs[0] if tf.paragraphs and tf.paragraphs[0].runs else None
    src_size = template_run.font.size if template_run else None
    src_name = template_run.font.name if template_run else None
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = link_text
    if src_size is not None:
        run.font.size = src_size
    if src_name is not None:
        run.font.name = src_name
    if url:
        try:
            run.hyperlink.address = url
        except Exception:
            pass


def _replace_picture(slide, picture_shape, image_bytes: bytes) -> None:
    """Replace the image data behind an existing PICTURE shape, keeping position.

    Adds the new image as a new image part on the slide, then redirects the
    picture's <a:blip r:embed> to the new relationship.
    """
    image_part, rId = slide.part.get_or_add_image(io.BytesIO(image_bytes))
    blip = picture_shape.element.find(f".//{{{A_NS}}}blip")
    if blip is not None:
        blip.set(f"{{{R_NS}}}embed", rId)


def _fill_slide(slide, case: Case) -> None:
    title_shape = _find_shape(slide, SHAPE_TITLE)
    if title_shape is not None:
        _set_single_text(title_shape, f"事例- {case.title} -")

    subtitle_shape = _find_shape(slide, SHAPE_SUBTITLE)
    if subtitle_shape is not None:
        _set_single_text(subtitle_shape, case.subtitle)

    overview_shape = _find_shape(slide, SHAPE_OVERVIEW)
    if overview_shape is not None:
        _set_text_preserving_format(overview_shape, case.overview or ["(情報なし)"])

    challenges_shape = _find_shape(slide, SHAPE_CHALLENGES)
    if challenges_shape is not None:
        _set_text_preserving_format(challenges_shape, case.challenges or ["(情報なし)"])

    solutions_shape = _find_shape(slide, SHAPE_SOLUTIONS)
    if solutions_shape is not None:
        _set_text_preserving_format(solutions_shape, case.solutions or ["(情報なし)"])

    link_shape = _find_shape(slide, SHAPE_LINK)
    if link_shape is not None:
        _set_link_textbox(slide, link_shape, case.link_text or case.title, case.url)

    if case.image_url:
        img = scraper.download_image(case.image_url)
        if img is not None:
            picture_shape = _find_shape(slide, SHAPE_PICTURE)
            if picture_shape is not None:
                _replace_picture(slide, picture_shape, img[0])


def build_pptx(cases: list[Case], template_path: Path = TEMPLATE_PATH) -> bytes:
    """Build a single .pptx containing one slide per case, based on the template.

    Clones the template slide within the same Presentation instance so that
    shared parts (layouts, theme, media) are not duplicated in the package.
    """
    if not cases:
        raise ValueError("cases must not be empty")

    prs = Presentation(str(template_path))
    src_slide = prs.slides[0]

    # Snapshot the pristine shape XML and rel map BEFORE we modify the first slide,
    # so subsequent slides can be cloned from the original state.
    original_shape_xmls = [deepcopy(shp.element) for shp in src_slide.shapes]
    original_rels = {rId: rel for rId, rel in src_slide.part.rels.items()}

    _fill_slide(src_slide, cases[0])

    rid_attrs = (f"{{{R_NS}}}embed", f"{{{R_NS}}}link", f"{{{R_NS}}}id")
    for case in cases[1:]:
        new_slide = prs.slides.add_slide(src_slide.slide_layout)
        for shp in list(new_slide.shapes):
            shp.element.getparent().remove(shp.element)

        rid_map: dict[str, str] = {}
        for src_rId, rel in original_rels.items():
            if rel.is_external:
                new_rId = new_slide.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
            rid_map[src_rId] = new_rId

        for shp_xml in original_shape_xmls:
            new_el = deepcopy(shp_xml)
            for el in new_el.iter():
                for attr in rid_attrs:
                    if attr in el.attrib:
                        old = el.attrib[attr]
                        if old in rid_map:
                            el.attrib[attr] = rid_map[old]
            new_slide.shapes._spTree.append(new_el)

        _fill_slide(new_slide, case)

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


def build_single_pptx(case: Case, template_path: Path = TEMPLATE_PATH) -> bytes:
    prs = Presentation(str(template_path))
    _fill_slide(prs.slides[0], case)
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()
