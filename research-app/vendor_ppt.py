"""Slide generator for vendor research, recreating the screenshot layout from scratch.

The slide structure mirrors the reference image:
    ┌──────────────────────────────────────────────────────┐
    │ {社名}の「{製品名}」は{summary}                       │
    ├─────┬──────────────────────────┬─────────────────────┤
    │ 出展者│ {company}                │                     │
    │ 製品名│ {product}                │                     │
    │ 製品の│ • feature 1              │   [ product photo ] │
    │ 特長 │ • feature 2              │                     │
    │      │ • feature 3              │                     │
    │ 解決 │ • problem 1              │                     │
    │ する │ • problem 2              │                     │
    │ 課題 │                          │                     │
    │ 活用 │ • 見出し: 説明           │                     │
    │ 例   │ • 見出し: 説明           │                     │
    └─────┴──────────────────────────┴─────────────────────┘

When the user provides a real .pptx template we will switch to template-based
generation (same approach as ppt_export.py).
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from models import VendorCase
import scraper

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

LABEL_FILL = RGBColor(0xDC, 0xDC, 0xDC)
LABEL_BORDER = RGBColor(0xA0, 0xA0, 0xA0)
ROW_BORDER = RGBColor(0xC0, 0xC0, 0xC0)

LABEL_X = Inches(0.4)
LABEL_W = Inches(1.1)
CONTENT_X = Inches(1.6)
CONTENT_W = Inches(5.6)
IMAGE_X = Inches(7.6)
IMAGE_Y = Inches(1.4)
IMAGE_W = Inches(5.4)
IMAGE_H = Inches(5.4)


def _add_title_band(slide, vendor: VendorCase) -> None:
    box = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = f"{vendor.company}の「{vendor.product}」は{vendor.summary}"
    run.font.size = Pt(14)
    run.font.bold = True

    # Section header "内容".
    hdr = slide.shapes.add_textbox(CONTENT_X, Inches(1.05), CONTENT_W, Inches(0.3))
    tf = hdr.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "内容"
    run.font.size = Pt(11)
    run.font.bold = True


def _add_label_box(slide, y: int, h: int, text: str):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LABEL_X, y, LABEL_W, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = LABEL_FILL
    shp.line.color.rgb = LABEL_BORDER
    shp.line.width = Pt(0.5)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    return shp


def _add_content_box(
    slide,
    y: int,
    h: int,
    lines: list[str],
    *,
    bulleted: bool = True,
    font_size_pt: int = 11,
) -> None:
    txbox = slide.shapes.add_textbox(CONTENT_X, y, CONTENT_W, h)
    tf = txbox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        prefix = "・ " if bulleted else ""
        # Support "見出し: 説明" with bold header (matches the screenshot's 活用例 style).
        if bulleted and (":" in line or "：" in line):
            sep = "：" if "：" in line else ":"
            header, _, rest = line.partition(sep)
            if rest.strip():
                head_run = p.add_run()
                head_run.text = f"{prefix}{header}{sep}"
                head_run.font.size = Pt(font_size_pt)
                head_run.font.bold = True
                body_run = p.add_run()
                body_run.text = rest.strip()
                body_run.font.size = Pt(font_size_pt)
                continue
        run = p.add_run()
        run.text = f"{prefix}{line}"
        run.font.size = Pt(font_size_pt)


def _add_image(slide, vendor: VendorCase) -> None:
    img = scraper.get_case_image(vendor.url, vendor.image_url)
    if img is None:
        return
    try:
        slide.shapes.add_picture(io.BytesIO(img[0]), IMAGE_X, IMAGE_Y, IMAGE_W, IMAGE_H)
    except Exception:
        pass


def _build_slide(prs, vendor: VendorCase) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    _add_title_band(slide, vendor)

    # 5 rows, with heights matching the screenshot proportions.
    rows = [
        ("出展者", [vendor.company], Inches(1.45), Inches(0.55), False),
        ("製品名", [vendor.product], Inches(2.05), Inches(0.55), False),
        ("製品の特長", vendor.features[:3] or ["(情報なし)"], Inches(2.65), Inches(1.65), True),
        ("解決する課題", vendor.problems_solved[:3] or ["(情報なし)"], Inches(4.35), Inches(1.05), True),
        ("活用例", vendor.use_cases[:3] or ["(情報なし)"], Inches(5.45), Inches(1.55), True),
    ]
    for label, content, y, h, bulleted in rows:
        _add_label_box(slide, y, h, label)
        _add_content_box(slide, y, h, content, bulleted=bulleted)

    _add_image(slide, vendor)


def build_vendor_pptx(vendors: list[VendorCase]) -> bytes:
    if not vendors:
        raise ValueError("vendors must not be empty")
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for v in vendors:
        _build_slide(prs, v)
    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()
