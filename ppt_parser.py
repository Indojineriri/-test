from __future__ import annotations

import base64
import glob
import io
import os
import subprocess
import tempfile
from dataclasses import dataclass


@dataclass
class Slide:
    index: int  # 1-based
    text: str
    image_b64: str | None = None
    image_media_type: str = "image/png"


def _extract_text(pptx_bytes: bytes) -> list[str]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_bytes))
    out: list[str] = []
    for slide in prs.slides:
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        # Speaker notes add useful intent for meeting prep.
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                parts.append(f"[ノート] {note}")
        out.append("\n".join(parts))
    return out


def _pptx_to_pdf(pptx_bytes: bytes, workdir: str) -> str:
    pptx_path = os.path.join(workdir, "deck.pptx")
    with open(pptx_path, "wb") as f:
        f.write(pptx_bytes)
    profile = os.path.join(workdir, "lo-profile")
    cmd = [
        "soffice",
        "--headless",
        "--norestore",
        f"-env:UserInstallation=file://{profile}",
        "--convert-to",
        "pdf",
        "--outdir",
        workdir,
        pptx_path,
    ]
    subprocess.run(cmd, check=True, capture_output=True, timeout=240)
    pdfs = glob.glob(os.path.join(workdir, "*.pdf"))
    if not pdfs:
        raise RuntimeError("LibreOffice failed to produce a PDF from the deck.")
    return pdfs[0]


def _render_pdf(pdf_path: str, long_edge: int) -> list[str]:
    import fitz  # PyMuPDF

    images: list[str] = []
    doc = fitz.open(pdf_path)
    try:
        for page in doc:
            rect = page.rect
            edge = max(rect.width, rect.height) or 1
            zoom = long_edge / edge
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
            png = pix.tobytes("png")
            images.append(base64.standard_b64encode(png).decode("utf-8"))
    finally:
        doc.close()
    return images


def parse_text(pptx_bytes: bytes) -> list[Slide]:
    """Extract per-slide text. Always succeeds for a valid .pptx."""
    return [Slide(index=i + 1, text=t) for i, t in enumerate(_extract_text(pptx_bytes))]


def render_images_into(pptx_bytes: bytes, slides: list[Slide], long_edge: int = 1600) -> None:
    """Rasterise each slide and attach the PNG to the matching Slide.

    Requires LibreOffice (the `soffice` binary with the Impress module). Raises
    on failure so the caller can fall back to text-only.
    """
    with tempfile.TemporaryDirectory() as workdir:
        pdf_path = _pptx_to_pdf(pptx_bytes, workdir)
        images = _render_pdf(pdf_path, long_edge)
    for i, slide in enumerate(slides):
        if i < len(images):
            slide.image_b64 = images[i]


def parse_deck(pptx_bytes: bytes, render_images: bool, long_edge: int = 1600) -> list[Slide]:
    """Convenience: text plus optional images in one call."""
    slides = parse_text(pptx_bytes)
    if render_images:
        render_images_into(pptx_bytes, slides, long_edge)
    return slides


def extract_minutes_text(filename: str, data: bytes) -> str:
    """Best-effort plain-text extraction from an optional minutes file."""
    ext = os.path.splitext(filename)[1].lower()
    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        import fitz

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            return "\n".join(page.get_text() for page in doc)
        finally:
            doc.close()
    if ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"対応していない議事録の形式です: {ext}")
