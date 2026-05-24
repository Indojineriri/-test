from __future__ import annotations

import io

from pptx import Presentation
from pptx.util import Pt

from claude_client import MeetingMaterial


def build_pptx(m: MeetingMaterial) -> bytes:
    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = m.meeting_title
    slide.placeholders[1].text = m.overall_objective

    # Agenda overview
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "アジェンダ"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, item in enumerate(m.agenda, 1):
        p = tf.paragraphs[0] if i == 1 else tf.add_paragraph()
        p.text = f"{i}. {item.title}（{item.duration_min}分）"

    # One slide per agenda item
    for i, item in enumerate(m.agenda, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{i}. {item.title}"
        tf = slide.placeholders[1].text_frame
        tf.clear()
        tf.paragraphs[0].text = f"ねらい: {item.objective}"
        for tp in item.talking_points:
            p = tf.add_paragraph()
            p.text = tp
            p.level = 1

    # Key messages
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "強調すべきメッセージ"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, km in enumerate(m.key_messages):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = km
        p.font.size = Pt(20)

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()
